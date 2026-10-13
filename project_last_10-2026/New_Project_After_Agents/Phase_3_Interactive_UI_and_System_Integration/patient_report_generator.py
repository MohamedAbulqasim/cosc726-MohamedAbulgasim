"""
Patient Diagnostic Word Report Generator (.docx)
Produces an official clinical summary report for the patient in Arabic or English.
"""

import os
import sys
import docx
from datetime import datetime

sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
import common_docx_styles as cds

sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'Phase_2_Multi_Agent_Architecture')))
from schemas import MultiAgentDiagnosticDossier

def generate_patient_docx(dossier: MultiAgentDiagnosticDossier, output_path: str = None, lang: str = "ar") -> str:
    """Generates an official patient diagnostic report (.docx)."""
    p_in = dossier.patient_input
    triage = dossier.triage
    pred = dossier.prediction
    reason = dossier.reasoning
    life = dossier.lifestyle
    gate = dossier.gatekeeper

    doc = docx.Document()
    is_ar = (lang == "ar")

    if is_ar:
        cds.make_document_rtl(doc)

    title_ar = f"تقرير التشخيص والتقييم السريري الشامل: {p_in.patient_name}"
    title_en = f"Comprehensive Clinical Diagnostic Report: {p_in.patient_name}"

    cds.add_header_banner(
        doc,
        title_ar=title_ar,
        title_en=title_en,
        doc_type="ملف تشخيصي صادر عن نظام الذكاء الاصطناعي متعدد الوكلاء" if is_ar else "Agentic Clinical Decision Support System Dossier",
        rtl=is_ar,
        include_academic_meta=False
    )

    if is_ar:
        # ARABIC REPORT
        cds.add_heading_ar(doc, "1. بيانات المريض والمؤشرات الحيوية الأساسية", 1)
        gender_text_ar = "ذكر" if str(getattr(p_in, 'gender', 'Female')).lower() == 'male' else "أنثى"
        patient_meta = [
            ["معرف المريض:", p_in.patient_id, "تاريخ الفحص:", datetime.now().strftime("%Y-%m-%d %H:%M")],
            ["نوع المريض (الجنس):", gender_text_ar, "العمر:", f"{p_in.age} سنة"],
            ["مؤشر كتلة الجسم (BMI):", f"{p_in.bmi:.1f} kg/m²", "مرات الحمل:", f"{p_in.pregnancies}" if gender_text_ar == "أنثى" else "غير متاح (ذكر)"],
            ["مستوى الجلوكوز:", f"{p_in.glucose:.1f} mg/dL", "ضغط الدم الانبساطي:", f"{p_in.blood_pressure:.1f} mmHg"]
        ]
        cds.add_styled_table(doc, ["البيان", "القيمة", "البيان", "القيمة"], patient_meta, rtl=True, col_widths=[1.6, 1.9, 1.6, 1.9])

        # Emergency Alert if triggered
        if gate.emergency_alert_triggered:
            cds.add_callout_box(
                doc,
                gate.emergency_message_ar or "تنبيه طوارئ حاد!",
                title="إنذار سريري حرج - عناية طبية فورية",
                box_type="alert",
                rtl=True
            )

        # Prediction Summary
        cds.add_heading_ar(doc, "2. نتيجة التشخيص التنبئي المعاير", 1)
        diag_color_box = "warning" if pred.prediction_label == 1 else "gate"
        clean_diag_ar = pred.diagnosis_ar.replace("أيضياً", "").strip()
        while "  " in clean_diag_ar:
            clean_diag_ar = clean_diag_ar.replace("  ", " ")
        cds.add_callout_box(
            doc,
            f"التشخيص التنبئي: {clean_diag_ar}\n"
            f"احتمالية الخطر المعايرة: {pred.diabetes_probability:.1f}%\n"
            f"مستوى الخطورة السريرية: {pred.risk_level_ar}\n"
            f"درجة الموثوقية الإحصائية: {pred.confidence_score:.1f}%",
            title="ملخص قرار الوكيل التنبؤي",
            box_type=diag_color_box,
            rtl=True
        )

        # Biomarkers Table
        cds.add_heading_ar(doc, "3. تحليل المؤشرات الحيوية ومساهمتها في الخطر", 1)
        bio_rows = []
        for bio in pred.contributing_biomarkers:
            bio_rows.append([
                bio.feature_name_ar,
                str(bio.patient_value),
                bio.ideal_range,
                "مرتفع" if bio.risk_impact == "Elevated" else "طبيعي"
            ])
        cds.add_styled_table(doc, ["المؤشر الحيوي", "قيمة المريض", "المدى الطبيعي", "تقييم الأثر"], bio_rows, rtl=True, col_widths=[2.2, 1.6, 1.8, 1.4])

        # Clinical Reasoning
        cds.add_heading_ar(doc, "4. التعليل السريري وإرشادات ADA", 1)
        cds.add_paragraph_ar(doc, reason.detailed_explanation_ar)
        if reason.key_drivers_identified:
            cds.add_paragraph_ar(doc, "أهم العوامل الدافعة للتشخيص:", bold=True)
            for driver in reason.key_drivers_identified:
                cds.add_paragraph_ar(doc, f" • {driver}")

        # Lifestyle & Nutrition
        cds.add_heading_ar(doc, "5. خطة التغذية ونمط الحياة المخصصة", 1)
        cds.add_paragraph_ar(doc, "التوصيات الغذائية:", bold=True)
        for d in life.dietary_recommendations_ar:
            cds.add_paragraph_ar(doc, f" • {d}")
        
        cds.add_paragraph_ar(doc, f"برنامج النشاط البدني: {life.physical_activity_plan_ar}")

        cds.add_paragraph_ar(doc, "جدول الفحوصات والمتابعة السريرية:", bold=True)
        for s in life.monitoring_schedule_ar:
            cds.add_paragraph_ar(doc, f" • {s}")

        # Gatekeeper Disclaimer
        cds.add_heading_ar(doc, "6. اعتماد بوابات الأمان وإخلاء المسؤولية", 1)
        cds.add_callout_box(
            doc,
            gate.safety_disclaimer_ar + f"\n\nسجل تدقيق الأمان: {', '.join(gate.audit_trail)}",
            title="اعتماد الحوكمة الطبية",
            box_type="info",
            rtl=True
        )

    else:
        # ENGLISH REPORT
        cds.add_heading_en(doc, "1. Patient Metadata & Baseline Biomarkers", 1)
        gender_text_en = "Male" if str(getattr(p_in, 'gender', 'Female')).lower() == 'male' else "Female"
        patient_meta_en = [
            ["Patient ID:", p_in.patient_id, "Date & Time:", datetime.now().strftime("%Y-%m-%d %H:%M")],
            ["Biological Sex / Gender:", gender_text_en, "Age:", f"{p_in.age} years"],
            ["BMI:", f"{p_in.bmi:.1f} kg/m²", "Pregnancies:", f"{p_in.pregnancies}" if gender_text_en == "Female" else "N/A (Male)"],
            ["Plasma Glucose:", f"{p_in.glucose:.1f} mg/dL", "Diastolic Blood Pressure:", f"{p_in.blood_pressure:.1f} mmHg"]
        ]
        cds.add_styled_table(doc, ["Parameter", "Value", "Parameter", "Value"], patient_meta_en, rtl=False, col_widths=[1.8, 1.7, 2.0, 1.5])

        if gate.emergency_alert_triggered:
            cds.add_callout_box(
                doc,
                gate.emergency_message_en or "Acute Emergency Alert!",
                title="CRITICAL CLINICAL ALERT - IMMEDIATE CARE REQUIRED",
                box_type="alert",
                rtl=False
            )

        cds.add_heading_en(doc, "2. Calibrated Diagnostic Assessment", 1)
        box_style = "warning" if pred.prediction_label == 1 else "gate"
        cds.add_callout_box(
            doc,
            f"Diagnostic Outcome: {pred.diagnosis_en}\n"
            f"Calibrated Risk Probability: {pred.diabetes_probability:.1f}%\n"
            f"Clinical Risk Tier: {pred.risk_level}\n"
            f"Confidence Score: {pred.confidence_score:.1f}%",
            title="Predictive Agent Summary",
            box_type=box_style,
            rtl=False
        )

        cds.add_heading_en(doc, "3. Biomarker Impact & Risk Contribution", 1)
        bio_en_rows = []
        for bio in pred.contributing_biomarkers:
            bio_en_rows.append([
                bio.feature_name,
                str(bio.patient_value),
                bio.ideal_range,
                bio.risk_impact
            ])
        cds.add_styled_table(doc, ["Biomarker", "Patient Value", "Standard Range", "Clinical Impact"], bio_en_rows, rtl=False, col_widths=[2.0, 1.5, 2.0, 1.5])

        cds.add_heading_en(doc, "4. Clinical Reasoning & ADA Alignment", 1)
        cds.add_paragraph_en(doc, reason.detailed_explanation_en)

        cds.add_heading_en(doc, "5. Personalized Lifestyle & Clinical Nutrition Plan", 1)
        cds.add_paragraph_en(doc, "Dietary Recommendations:", bold=True)
        for d in life.dietary_recommendations_en:
            cds.add_paragraph_en(doc, f" • {d}")
        
        cds.add_paragraph_en(doc, f"Physical Conditioning: {life.physical_activity_plan_en}")

        cds.add_paragraph_en(doc, "Monitoring & Diagnostic Schedule:", bold=True)
        for s in life.monitoring_schedule_en:
            cds.add_paragraph_en(doc, f" • {s}")

        cds.add_heading_en(doc, "6. Quality Governance & Medical Disclaimer", 1)
        cds.add_callout_box(
            doc,
            gate.safety_disclaimer_en + f"\n\nAudit Trail: {', '.join(gate.audit_trail)}",
            title="Clinical Safety Sign-Off",
            box_type="info",
            rtl=False
        )

    if not output_path:
        out_dir = os.path.dirname(os.path.abspath(__file__))
        output_path = os.path.join(out_dir, f"Patient_Report_{p_in.patient_id}_{lang}.docx")

    doc.save(output_path)
    return output_path
