"""
Common Docx Styling and Helper Utilities for Neelain University Master's Project.
Supports:
- Professional Arabic (RTL) and English (LTR) typography
- University Header & Footer styling
- Themed tables with custom cell shading and subtle borders
- Highlighted Callout Boxes (Quality Gates, Clinical Alerts, Important Notes)
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

# Official Theme Colors
PRIMARY_COLOR = RGBColor(26, 77, 46)      # #1a4d2e Dark Clinical Green
SECONDARY_COLOR = RGBColor(30, 58, 138)   # #1e3a8a Deep Royal Blue
ACCENT_COLOR = RGBColor(180, 83, 9)       # #b45309 Amber/Warning
TEXT_COLOR = RGBColor(30, 41, 59)         # #1e293b Slate 800
MUTED_COLOR = RGBColor(100, 116, 139)     # #64748b Slate 500
BORDER_COLOR_HEX = "CBD5E1"
HEADER_BG_HEX = "1A4D2E"
HEADER_TEXT_HEX = "FFFFFF"
ALT_ROW_BG_HEX = "F8FAFC"
CALLOUT_BG_HEX = "F0FDF4"
CALLOUT_BORDER_HEX = "16A34A"

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for a cell in twips."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def format_run_rtl(run, font_name="Traditional Arabic"):
    """Enforces RTL text direction and complex script font on a run."""
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    if not rPr.xpath('w:rtl'):
        rPr.append(parse_xml(f'<w:rtl {nsdecls("w")}/>'))
    rFonts = rPr.xpath('w:rFonts')
    if rFonts:
        rFonts[0].set(qn('w:cs'), font_name)
    else:
        rPr.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:cs="{font_name}"/>'))

def make_document_rtl(doc):
    """Enables document-wide RTL section layout."""
    for section in doc.sections:
        sectPr = section._sectPr
        if not sectPr.xpath('w:bidi'):
            sectPr.append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))

def set_paragraph_direction(p, rtl=False):
    """Sets text direction to RTL or LTR."""
    pPr = p._element.get_or_add_pPr()
    if rtl:
        if not pPr.xpath('w:bidi'):
            pPr.append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

def apply_table_borders(table, border_hex="CBD5E1"):
    """Applies subtle borders to a docx table."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_header_banner(doc, title_ar, title_en, doc_type="وثيقة مشروع ماجستير / Master's Project Document", rtl=False, include_academic_meta=True):
    """Adds formal Neelain University document header with native RTL support."""
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        if rtl:
            sectPr = section._sectPr
            if not sectPr.xpath('w:bidi'):
                sectPr.append(parse_xml(f'<w:bidi {nsdecls("w")}/>'))

    # University Header Table
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    if rtl:
        tblPr = tbl._element.xpath('w:tblPr')
        if tblPr and not tblPr[0].xpath('w:bidiVisual'):
            tblPr[0].append(parse_xml(f'<w:bidiVisual {nsdecls("w")}/>'))

    # If rtl=True with bidiVisual, Cell 0 is on the RIGHT and Cell 1 on the LEFT.
    # So Cell 0 gets Arabic metadata (Right) and Cell 1 gets English metadata (Left).
    if rtl:
        cell_ar = tbl.rows[0].cells[0]
        cell_en = tbl.rows[0].cells[1]
    else:
        cell_en = tbl.rows[0].cells[0]
        cell_ar = tbl.rows[0].cells[1]

    cell_en.width = Inches(3.5)
    cell_ar.width = Inches(3.5)

    # English metadata
    p_en = cell_en.paragraphs[0]
    p_en.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1 = p_en.add_run("Al-Neelain University\nFaculty of Graduate Studies\nMaster's in Information Technology / CS")
    r1.font.name = "Arial"
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = MUTED_COLOR
    r1.font.bold = True

    # Arabic metadata
    p_ar = cell_ar.paragraphs[0]
    set_paragraph_direction(p_ar, rtl=True)
    r2 = p_ar.add_run("جامعة النيلين\nكلية الدراسات العليا\nماجستير تقنية المعلومات / علوم الحاسوب")
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = MUTED_COLOR
    r2.font.bold = True
    format_run_rtl(r2, "Traditional Arabic")

    set_cell_margins(cell_en, top=60, bottom=60, left=60, right=60)
    set_cell_margins(cell_ar, top=60, bottom=60, left=60, right=60)

    # Dividing Line
    doc.add_paragraph()
    
    # Main Title Box
    title_tbl = doc.add_table(rows=1, cols=1)
    title_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    if rtl:
        ttblPr = title_tbl._element.xpath('w:tblPr')
        if ttblPr and not ttblPr[0].xpath('w:bidiVisual'):
            ttblPr[0].append(parse_xml(f'<w:bidiVisual {nsdecls("w")}/>'))

    t_cell = title_tbl.rows[0].cells[0]
    t_cell.width = Inches(7.0)
    set_cell_background(t_cell, "1A4D2E")
    set_cell_margins(t_cell, top=160, bottom=160, left=180, right=180)

    tp1 = t_cell.paragraphs[0]
    tp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if rtl:
        set_paragraph_direction(tp1, rtl=True)
        tp1.alignment = WD_ALIGN_PARAGRAPH.CENTER

    tr1 = tp1.add_run(title_ar + "\n")
    tr1.font.size = Pt(16)
    tr1.font.bold = True
    tr1.font.color.rgb = RGBColor(255, 255, 255)
    format_run_rtl(tr1, "Traditional Arabic")

    tr2 = tp1.add_run(title_en + "\n")
    tr2.font.name = "Calibri"
    tr2.font.size = Pt(12)
    tr2.font.bold = True
    tr2.font.color.rgb = RGBColor(222, 255, 154)

    tr3 = tp1.add_run(f"[{doc_type}]")
    tr3.font.name = "Calibri"
    tr3.font.size = Pt(9.5)
    tr3.font.color.rgb = RGBColor(255, 255, 255)

    # Author & Academic metadata (Only for academic project reports, excluded for patient diagnostic reports)
    if include_academic_meta:
        doc.add_paragraph()
        meta_p = doc.add_paragraph()
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if rtl:
            set_paragraph_direction(meta_p, rtl=True)
            meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            m1 = meta_p.add_run("إعداد الباحث: ")
            m1.font.bold = True
            m1.font.size = Pt(10)
            format_run_rtl(m1, "Traditional Arabic")

            m2 = meta_p.add_run("محمد ابوالقاسم  |  ")
            m2.font.bold = True
            m2.font.color.rgb = PRIMARY_COLOR
            m2.font.size = Pt(10)
            format_run_rtl(m2, "Traditional Arabic")

            m3 = meta_p.add_run("إشراف: ")
            m3.font.bold = True
            m3.font.size = Pt(10)
            format_run_rtl(m3, "Traditional Arabic")

            m4 = meta_p.add_run("د. فخرالدين سعيد  |  ")
            m4.font.bold = True
            m4.font.color.rgb = PRIMARY_COLOR
            m4.font.size = Pt(10)
            format_run_rtl(m4, "Traditional Arabic")

            m5 = meta_p.add_run("مادة ( course ) : COSC726 AGENTIC AI  |  ")
            m5.font.bold = True
            m5.font.color.rgb = SECONDARY_COLOR
            m5.font.size = Pt(10)
            format_run_rtl(m5, "Traditional Arabic")

            m6 = meta_p.add_run("التاريخ: 2026 م")
            m6.font.bold = True
            m6.font.size = Pt(10)
            format_run_rtl(m6, "Traditional Arabic")
        else:
            set_paragraph_direction(meta_p, rtl=False)
            meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            m1 = meta_p.add_run("Researcher: ")
            m1.font.name = "Calibri"
            m1.font.bold = True
            m1.font.size = Pt(9.5)

            m2 = meta_p.add_run("Mohamed Abuelgasim  |  ")
            m2.font.name = "Calibri"
            m2.font.bold = True
            m2.font.color.rgb = PRIMARY_COLOR
            m2.font.size = Pt(9.5)

            m3 = meta_p.add_run("Supervisor: ")
            m3.font.name = "Calibri"
            m3.font.bold = True
            m3.font.size = Pt(9.5)

            m4 = meta_p.add_run("Dr. Fakhreldeen Saeed  |  ")
            m4.font.name = "Calibri"
            m4.font.bold = True
            m4.font.color.rgb = PRIMARY_COLOR
            m4.font.size = Pt(9.5)

            m5 = meta_p.add_run("Course: COSC726 AGENTIC AI  |  ")
            m5.font.name = "Calibri"
            m5.font.bold = True
            m5.font.color.rgb = SECONDARY_COLOR
            m5.font.size = Pt(9.5)

            m6 = meta_p.add_run("Date: 2026")
            m6.font.name = "Calibri"
            m6.font.bold = True
            m6.font.size = Pt(9.5)

def add_heading_ar(doc, text, level=1):
    """Adds a styled Arabic heading with strict RTL direction."""
    p = doc.add_paragraph()
    set_paragraph_direction(p, rtl=True)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.bold = True
    format_run_rtl(run, "Traditional Arabic")
    if level == 1:
        run.font.size = Pt(15)
        run.font.color.rgb = PRIMARY_COLOR
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = SECONDARY_COLOR
    else:
        run.font.size = Pt(11.5)
        run.font.color.rgb = TEXT_COLOR
    return p

def add_heading_en(doc, text, level=1):
    """Adds a styled English heading."""
    p = doc.add_paragraph()
    set_paragraph_direction(p, rtl=False)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(15)
        run.font.color.rgb = PRIMARY_COLOR
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = SECONDARY_COLOR
    else:
        run.font.size = Pt(11.5)
        run.font.color.rgb = TEXT_COLOR
    return p

def add_paragraph_ar(doc, text, bold=False):
    """Adds an Arabic paragraph with strict RTL direction."""
    p = doc.add_paragraph()
    set_paragraph_direction(p, rtl=True)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.bold = bold
    run.font.color.rgb = TEXT_COLOR
    format_run_rtl(run, "Traditional Arabic")
    return p

def add_paragraph_en(doc, text, bold=False):
    """Adds an English paragraph."""
    p = doc.add_paragraph()
    set_paragraph_direction(p, rtl=False)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.font.bold = bold
    run.font.color.rgb = TEXT_COLOR
    return p

def add_callout_box(doc, text, title=None, box_type="info", rtl=False):
    """Adds a highlighted callout box for Quality Gates or Clinical Alerts with RTL support."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    if rtl:
        tblPr = tbl._element.xpath('w:tblPr')
        if tblPr and not tblPr[0].xpath('w:bidiVisual'):
            tblPr[0].append(parse_xml(f'<w:bidiVisual {nsdecls("w")}/>'))

    cell = tbl.rows[0].cells[0]
    cell.width = Inches(7.0)

    # Pick colors based on box_type
    if box_type == "gate":
        bg_hex = "F0FDF4"       # Light green
        border_hex = "16A34A"   # Green
        title_color = RGBColor(22, 101, 52)
    elif box_type == "alert":
        bg_hex = "FEF2F2"       # Light red
        border_hex = "DC2626"   # Red
        title_color = RGBColor(153, 27, 27)
    elif box_type == "warning":
        bg_hex = "FFFBEB"       # Light yellow/amber
        border_hex = "D97706"   # Amber
        title_color = RGBColor(146, 64, 14)
    else:
        bg_hex = "EFF6FF"       # Light blue
        border_hex = "2563EB"   # Blue
        title_color = RGBColor(30, 64, 175)

    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)

    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="{"none" if rtl else "single"}" w:sz="{"0" if rtl else "24"}" w:space="0" w:color="{border_hex}"/>'
        f'<w:right w:val="{"single" if rtl else "none"}" w:sz="{"24" if rtl else "0"}" w:space="0" w:color="{border_hex}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    set_paragraph_direction(p, rtl=rtl)
    
    if title:
        rt = p.add_run(f"[{title}]\n" if not rtl else f"【 {title} 】\n")
        rt.font.bold = True
        rt.font.size = Pt(11)
        rt.font.color.rgb = title_color
        if rtl:
            format_run_rtl(rt, "Traditional Arabic")
        else:
            rt.font.name = "Calibri"

    rc = p.add_run(text)
    rc.font.size = Pt(10)
    rc.font.color.rgb = TEXT_COLOR
    if rtl:
        format_run_rtl(rc, "Traditional Arabic")
    else:
        rc.font.name = "Calibri"

    doc.add_paragraph()  # Spacing

def add_styled_table(doc, headers, rows_data, rtl=False, col_widths=None):
    """Creates a beautifully styled table with shaded headers and alternating rows with RTL support."""
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_table_borders(tbl)

    if rtl:
        tblPr = tbl._element.xpath('w:tblPr')
        if tblPr and not tblPr[0].xpath('w:bidiVisual'):
            tblPr[0].append(parse_xml(f'<w:bidiVisual {nsdecls("w")}/>'))

    # Header Row
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        set_paragraph_direction(p, rtl=rtl)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)
        if rtl:
            format_run_rtl(r, "Traditional Arabic")
        else:
            r.font.name = "Calibri"
        set_cell_background(hdr_cells[i], HEADER_BG_HEX)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)

    # Data Rows
    for row_idx, data_row in enumerate(rows_data):
        row_cells = tbl.rows[row_idx + 1].cells
        bg_color = ALT_ROW_BG_HEX if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(data_row):
            row_cells[col_idx].text = ""
            p = row_cells[col_idx].paragraphs[0]
            set_paragraph_direction(p, rtl=rtl)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(cell_value))
            r.font.size = Pt(9.5)
            r.font.color.rgb = TEXT_COLOR
            if rtl:
                format_run_rtl(r, "Traditional Arabic")
            else:
                r.font.name = "Calibri"
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=120, right=120)

    # Column widths
    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Inches(width)

    doc.add_paragraph()  # Spacing
    return tbl
